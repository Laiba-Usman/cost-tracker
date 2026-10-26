"""Downloadable reports: PDF (with charts and the AI analysis) and Excel (with live formulas)."""
import io
from datetime import date

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, KeepTogether
from xml.sax.saxutils import escape

import calc

INK, OVER, UNDER, GREY = "#1f3a5f", "#a4332b", "#2f6f5e", "#8a9099"


def _fmt(x, cur):
    return "-" if x is None else "%s %s" % (cur, format(round(x), ",d"))


def _fig_png(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=170, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def make_charts(result):
    data = calc.chart_data(result)
    pngs = []
    c1 = data["by_category"]
    fig, ax = plt.subplots(figsize=(7, 3.2))
    if c1["labels"]:
        import numpy as np
        x = np.arange(len(c1["labels"]))
        ax.bar(x - 0.2, c1["estimated"], 0.4, label="Estimated", color="#b9c4d4")
        ax.bar(x + 0.2, c1["actual"], 0.4, label="Actual", color=INK)
        ax.set_xticks(x)
        ax.set_xticklabels([l if len(l) < 18 else l[:16] + "..." for l in c1["labels"]], rotation=20, ha="right", fontsize=8)
        ax.legend(frameon=False, fontsize=8)
    ax.set_title("Estimated vs actual cost by category (items with an actual)", fontsize=9, loc="left")
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(axis="y", labelsize=8)
    pngs.append(_fig_png(fig))

    c2 = data["top_variance"]
    fig, ax = plt.subplots(figsize=(7, 3.2))
    if c2["labels"]:
        labs = [l if len(l) < 38 else l[:36] + "..." for l in c2["labels"]][::-1]
        vals = c2["pct"][::-1]
        ax.barh(labs, vals, color=[OVER if v > 0 else UNDER for v in vals])
        ax.axvline(0, color="#444", lw=0.6)
        ax.tick_params(axis="y", labelsize=7)
    ax.set_title("Largest % variance by item (red = over, green = under)", fontsize=9, loc="left")
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(axis="x", labelsize=8)
    pngs.append(_fig_png(fig))

    c3 = data["actual_by_type"]
    fig, ax = plt.subplots(figsize=(4.2, 3.0))
    if c3["labels"]:
        tot = sum(c3["actual"])
        wedges, _ = ax.pie(c3["actual"], startangle=90,
               colors=["#1f3a5f", "#7d93b2", "#c9d2de", "#e4e7eb"][:len(c3["labels"])],
               wedgeprops={"width": 0.45, "edgecolor": "white"})
        ax.legend(wedges, ["%s: %s (%.1f%%)" % (l, format(round(v), ",d"), v / tot * 100) for l, v in zip(c3["labels"], c3["actual"])],
                  loc="center left", bbox_to_anchor=(1, 0.5), frameon=False, fontsize=8)
    ax.set_title("Actual cost by cost type", fontsize=9)
    pngs.append(_fig_png(fig))
    return pngs


def build_pdf(project, result, analysis):
    cur = project.get("currency") or "Rs."
    s = result["summary"]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title="Cost report - " + project["name"], author="Cost Tracker")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontSize=18, alignment=0, textColor=colors.HexColor(INK), spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12, textColor=colors.HexColor(INK), spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=9, leading=12.5)
    small = ParagraphStyle("s", parent=body, fontSize=7.5, leading=10, textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("c", parent=body, fontSize=7.5, leading=9)

    el = [Paragraph(escape(project["name"]), h1),
          Paragraph("Cost report generated on %s. Variance = actual - estimated; positive means overspent." % date.today().strftime("%d %b %Y"), small),
          Spacer(1, 6)]

    pct = "n/a" if s["pct"] is None else "%+.1f%%" % s["pct"]
    kp = [["Estimated (all items)", "Actual (items with actual)", "Variance", "% variance", "Forecast at completion"],
          [_fmt(s["estimated_total"], cur), _fmt(s["actual_total"], cur), _fmt(s["variance"], cur), pct, _fmt(s["forecast"], cur)]]
    t = Table(kp, colWidths=[36 * mm] * 5)
    t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, 0), 7), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#555555")),
                           ("FONTSIZE", (0, 1), (-1, 1), 10), ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                           ("LINEBELOW", (0, 1), (-1, 1), 0.5, colors.HexColor("#cccccc")),
                           ("LINEABOVE", (0, 0), (-1, 0), 0.5, colors.HexColor("#cccccc"))]))
    el += [t, Spacer(1, 4),
           Paragraph("%d of %d items have an actual cost. Contingency: %.0f%% (%s). Approved budget with contingency: %s." % (
               s["items_with_actual"], s["items"], s["contingency_pct"], _fmt(s["contingency_amount"], cur), _fmt(s["approved_budget"], cur)), body)]

    el.append(Paragraph("Cost analysis (AI-assisted)", h2))
    for para in analysis["summary"]:
        el += [Paragraph(escape(para), body), Spacer(1, 3)]
    if analysis["flags"]:
        rows = [["Item", "Estimated", "Actual", "% var.", "Why it was flagged"]]
        for f in analysis["flags"]:
            rows.append([Paragraph(escape(f["item"]), cell), _fmt(f["estimated"], cur), _fmt(f["actual"], cur),
                         "n/a" if f["pct"] is None else "%+.1f%%" % f["pct"], Paragraph(escape(" ".join(f["reasons"])), cell)])
        ft = Table(rows, colWidths=[45 * mm, 20 * mm, 20 * mm, 14 * mm, 75 * mm], repeatRows=1)
        ft.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7.5), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        el += [Spacer(1, 4), Paragraph("Unusual expenses", h2), ft]
    el.append(Paragraph("Method: " + escape(analysis["method"]), small))

    c1, c2, c3 = make_charts(result)
    el.append(KeepTogether([Paragraph("Charts", h2), Image(c1, width=170 * mm, height=170 * mm * 0.46)]))
    el.append(Image(c2, width=170 * mm, height=170 * mm * 0.46))
    el.append(Image(c3, width=95 * mm, height=95 * mm * 0.72))

    el.append(Paragraph("Category summary", h2))
    rows = [["Category", "Estimated", "Est. (with actual)", "Actual", "Variance", "% var."]]
    for c in result["categories"]:
        rows.append([Paragraph(escape(c["name"]), cell), _fmt(c["estimated"], cur), _fmt(c["est_compared"], cur), _fmt(c["actual"], cur),
                     _fmt(c["variance"], cur), "n/a" if c["pct"] is None else "%+.1f%%" % c["pct"]])
    ct = Table(rows, colWidths=[52 * mm, 24 * mm, 26 * mm, 24 * mm, 24 * mm, 14 * mm], repeatRows=1)
    ct.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7.5), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                            ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc"))]))
    el.append(ct)

    el.append(Paragraph("All items", h2))
    rows = [["Category", "Item", "Type", "Estimated", "Actual", "Variance", "% var.", "Status"]]
    for r in result["rows"]:
        rows.append([Paragraph(escape(r["category"]), cell), Paragraph(escape(r["item"]), cell), r["cost_type"],
                     _fmt(r["estimated"], cur), _fmt(r["actual"], cur), _fmt(r["variance"], cur),
                     "-" if r["pct"] is None else "%+.1f%%" % r["pct"], r["status"]])
    it = Table(rows, colWidths=[26 * mm, 46 * mm, 13 * mm, 20 * mm, 20 * mm, 20 * mm, 12 * mm, 20 * mm], repeatRows=1)
    it.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                            ("ALIGN", (3, 0), (6, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc"))]))
    el.append(it)

    def footer(canvas, d):
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(16 * mm, 9 * mm, "Cost report - " + project["name"][:70])
        canvas.drawRightString(A4[0] - 16 * mm, 9 * mm, "Page %d" % d.page)

    doc.build(el, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _text(ws, ref, value):
    """Write user text so Excel never treats it as a formula."""
    c = ws[ref]
    c.value = value
    c.data_type = "s"
    return c


def build_xlsx(project, result):
    cur = project.get("currency") or "Rs."
    wb = Workbook()
    ws = wb.active
    ws.title = "Cost data"
    head = ["Category", "Item", "Cost type", "Estimated (%s)" % cur, "Actual (%s)" % cur, "Variance (%s)" % cur, "% variance", "Status"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, r in enumerate(result["rows"], start=2):
        _text(ws, "A%d" % i, r["category"])
        _text(ws, "B%d" % i, r["item"])
        _text(ws, "C%d" % i, r["cost_type"])
        ws["D%d" % i] = r["estimated"]
        ws["E%d" % i] = r["actual"]
        ws["F%d" % i] = '=IF(E%d="","",E%d-D%d)' % (i, i, i)
        ws["G%d" % i] = '=IF(OR(E%d="",D%d=0),"",F%d/D%d)' % (i, i, i, i)
        ws["H%d" % i] = ('=IF(E{r}="","Not spent yet",IF(D{r}=0,IF(E{r}>0,"Unbudgeted","On track"),'
                         'IF(G{r}>0.05,"Over budget",IF(G{r}<-0.05,"Under budget","On track"))))').format(r=i)
        for col in "DEF":
            ws["%s%d" % (col, i)].number_format = "#,##0"
        ws["G%d" % i].number_format = "0.0%"
    last = len(result["rows"]) + 1
    for col, w in zip("ABCDEFGH", (30, 48, 11, 16, 16, 16, 12, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    sm = wb.create_sheet("Summary")
    rng = lambda col: "'Cost data'!$%s$2:$%s$%d" % (col, col, last)
    lines = [
        ("Project", project["name"], None),
        ("Estimated, all items", "=SUM(%s)" % rng("D"), "#,##0"),
        ("Actual, items with an actual", "=SUM(%s)" % rng("E"), "#,##0"),
        ("Estimated, items with an actual", '=SUMIF(%s,"<>",%s)' % (rng("E"), rng("D")), "#,##0"),
        ("Variance (actual - estimated)", "=B4-B5", "#,##0"),
        ("% variance", '=IF(B5=0,"",B6/B5)', "0.0%"),
        ("Not yet spent (estimate)", "=B3-B5", "#,##0"),
        ("Forecast at completion", "=B4+B8", "#,##0"),
        ("Contingency %", project["contingency_pct"] / 100.0, "0%"),
        ("Approved budget (estimate + contingency)", "=B3*(1+B10)", "#,##0"),
    ]
    for r, (a, b, fmt) in enumerate(lines, start=2):
        _text(sm, "A%d" % r, a)
        if r == 2:
            _text(sm, "B%d" % r, b)
        else:
            sm["B%d" % r] = b
            if fmt:
                sm["B%d" % r].number_format = fmt
    sm.column_dimensions["A"].width = 42
    sm.column_dimensions["B"].width = 28
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
