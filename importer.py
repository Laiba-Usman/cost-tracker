"""Read cost data from a CSV or Excel (.xlsx) file.

Two layouts are understood:
  1. A flat table with columns such as Category, Item, Cost type, Estimated, Actual.
  2. The LYKOS budget layout: a category name on its own row, items below it,
     "Subtotal ..." rows in between, and the estimate in a "Total (Rs.)" column.
Files are checked before they are parsed (size, real file type, zip size, row limit).
"""
import csv
import io
import math
import re
import zipfile

MAX_BYTES = 2 * 1024 * 1024
MAX_UNZIPPED = 20 * 1024 * 1024
MAX_ROWS = 1000
MAX_TEXT = 120

ITEM_H = {"item", "items", "description", "name", "expense", "expense item", "cost item"}
CAT_H = {"category", "cost category", "group"}
EST_H = {"estimated", "estimate", "estimated cost", "planned", "planned cost", "budget", "budgeted"}
EST_FALLBACK = {"total", "total cost", "amount", "cost"}
ACT_H = {"actual", "actual cost", "actuals", "spent"}
TYPE_H = {"cost type", "type"}


class ImportError_(ValueError):
    """Raised with a message that is safe to show to the user."""


def _norm(v):
    s = re.sub(r"\(.*?\)", "", str(v or "")).strip().lower()
    return re.sub(r"\s+", " ", s)


def to_number(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))
    if s in ("", "-", ".", "-."):
        return None
    try:
        f = float(s)
    except ValueError:
        return None
    return f if math.isfinite(f) else None


def _clean_text(v, limit=MAX_TEXT):
    s = re.sub(r"[\x00-\x1f\x7f]", " ", str(v or "")).strip()
    return s[:limit]


def _cost_type(v):
    s = _norm(v)
    return {"cash": "Cash", "in-kind": "In-kind", "in kind": "In-kind",
            "inkind": "In-kind", "free": "Free"}.get(s, "Other")


def check_upload(filename, data):
    """Validate the uploaded file. Returns 'csv' or 'xlsx'."""
    name = (filename or "").lower()
    if not (name.endswith(".csv") or name.endswith(".xlsx")):
        raise ImportError_("Only .csv and .xlsx files are accepted.")
    if not data:
        raise ImportError_("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ImportError_("The file is larger than 2 MB.")
    if name.endswith(".xlsx"):
        if not data.startswith(b"PK\x03\x04"):
            raise ImportError_("This file is not a real .xlsx workbook.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
                if "xl/workbook.xml" not in names:
                    raise ImportError_("This file is not a real .xlsx workbook.")
                if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                    raise ImportError_("The workbook is too large when unpacked.")
        except zipfile.BadZipFile:
            raise ImportError_("This file is not a real .xlsx workbook.")
        return "xlsx"
    if b"\x00" in data[:4096]:
        raise ImportError_("This does not look like a text CSV file.")
    return "csv"


def read_rows(kind, data):
    if kind == "csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        rows = []
        for i, r in enumerate(csv.reader(io.StringIO(text))):
            if i > MAX_ROWS + 20:
                raise ImportError_("Too many rows (limit is %d)." % MAX_ROWS)
            rows.append(r)
        return [rows]
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise ImportError_("The workbook could not be read.")
    sheets = []
    for ws in wb.worksheets[:10]:
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i > MAX_ROWS + 20:
                raise ImportError_("Too many rows (limit is %d)." % MAX_ROWS)
            rows.append(list(r))
        sheets.append(rows)
    return sheets


def _find_header(rows, need_estimate):
    for idx, row in enumerate(rows[:15]):
        cols = {}
        for c, v in enumerate(row):
            n = _norm(v)
            if not n:
                continue
            if n in ITEM_H: cols.setdefault("item", c)
            elif n in CAT_H: cols.setdefault("category", c)
            elif n in EST_H: cols.setdefault("est", c)
            elif n in EST_FALLBACK: cols.setdefault("est_fb", c)
            elif n in ACT_H: cols.setdefault("actual", c)
            elif n in TYPE_H: cols.setdefault("type", c)
        if "est" not in cols and "est_fb" in cols:
            cols["est"] = cols["est_fb"]
        if "item" in cols and (("est" in cols) if need_estimate else ("actual" in cols)):
            return idx, cols
    return None, None


def parse(sheets, need_estimate=True):
    """Return (items, warnings). Each item: category, item, cost_type, estimated, actual."""
    header = None
    for rows in sheets:
        idx, cols = _find_header(rows, need_estimate)
        if cols:
            header = (rows, idx, cols)
            break
    if not header:
        need = "Item and Estimated (or Total)" if need_estimate else "Item and Actual"
        raise ImportError_("No header row found. The file needs columns named " + need + ".")
    rows, idx, cols = header
    items, warnings = [], []
    current_cat = "Uncategorised"

    def cell(row, key):
        c = cols.get(key)
        return row[c] if c is not None and c < len(row) else None

    for n, row in enumerate(rows[idx + 1:], start=idx + 2):
        name = _clean_text(cell(row, "item"))
        if not name:
            continue
        low = name.lower()
        if low.startswith("subtotal") or low.startswith("total"):
            continue
        est = to_number(cell(row, "est"))
        act = to_number(cell(row, "actual"))
        others = [v for c, v in enumerate(row) if c != cols["item"] and v not in (None, "")]
        if "category" not in cols and est is None and act is None and not others:
            current_cat = name  # a category heading row
            continue
        if need_estimate and est is None:
            warnings.append("Row %d (%s) skipped: no estimated cost. If it is an Excel formula, open the file in Excel and save it first." % (n, name))
            continue
        if (est is not None and est < 0) or (act is not None and act < 0):
            warnings.append("Row %d (%s) skipped: negative amount." % (n, name))
            continue
        if (est or 0) > 1e10 or (act or 0) > 1e10:
            warnings.append("Row %d (%s) skipped: amount is too large." % (n, name))
            continue
        cat = _clean_text(cell(row, "category"), 60) if "category" in cols else current_cat
        items.append({"category": cat or "Uncategorised", "item": name,
                      "cost_type": _cost_type(cell(row, "type")),
                      "estimated": est if est is not None else 0.0, "actual": act})
        if len(items) > MAX_ROWS:
            raise ImportError_("Too many items (limit is %d)." % MAX_ROWS)
    if not items:
        raise ImportError_("No cost rows were found in the file.")
    return items, warnings


def parse_file(filename, data, need_estimate=True):
    kind = check_upload(filename, data)
    return parse(read_rows(kind, data), need_estimate)
